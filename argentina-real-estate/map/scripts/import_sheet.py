"""Convert the "Oportunidades Zona Norte" xlsx into the JSON files the site reads.

    python3 scripts/import_sheet.py path/to/Oportunidades_Zona_Norte.xlsx

Writes site/data/oportunidades.json (mapped) and site/data/avisos.json (table only).
Coordinates start at the barrio centroid with a small deterministic offset and
approx=true; scripts/geocode.mjs refines them at build time.
"""
import hashlib
import json
import math
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

SITE_DATA = Path(__file__).resolve().parent.parent / "site" / "data"

# Centroids for barrios/localidades outside caba.json.
EXTRA_CENTROIDS = {
    "Villa Ortuzar": (-34.5800, -58.4680),
    "Vicente Lopez": (-34.5260, -58.4760),
    "Florida": (-34.5280, -58.4930),
    "Olivos": (-34.5090, -58.4890),
    "La Lucila": (-34.4990, -58.4880),
    "Martinez": (-34.4920, -58.5100),
    "Acassuso": (-34.4780, -58.5040),
    "San Isidro": (-34.4720, -58.5280),
    "Beccar": (-34.4630, -58.5310),
}

OPORTUNIDADES_COLS = [
    "id", "zona", "barrio", "tipo", "direccion", "precio", "m2_cub", "m2_tot", "amb", "estado",
    "usd_m2", "ref_usd_m2", "bajo_ref", "obra_vs_ref", "valor_arq", "riesgos", "url", "portal",
    "clave_ref", "titulo", "alerta", "notas", "verificacion",
]
AVISOS_COLS = [
    "zona", "barrio", "tipo", "direccion", "precio", "m2_cub", "m2_tot", "amb", "estado",
    "usd_m2", "ref_usd_m2", "bajo_ref", "valor_arq", "riesgos", "url", "portal", "clave_ref",
]


def norm(s):
    return unicodedata.normalize("NFD", s or "").encode("ascii", "ignore").decode().strip().lower()


def centroids():
    caba = json.loads((SITE_DATA / "caba.json").read_text())
    out = {norm(b["name"]): (b["lat"], b["lon"]) for b in caba["barrios"]}
    out.update({norm(k): v for k, v in EXTRA_CENTROIDS.items()})
    return out


def jitter(key, lat, lon, meters=350):
    # Deterministic offset so markers in the same barrio don't stack.
    h = hashlib.sha1(key.encode()).digest()
    angle = h[0] / 255 * 2 * math.pi
    dist = (0.3 + 0.7 * h[1] / 255) * meters
    dlat = dist * math.cos(angle) / 111_320
    dlon = dist * math.sin(angle) / (111_320 * math.cos(math.radians(lat)))
    return round(lat + dlat, 6), round(lon + dlon, 6)


def clean(v):
    if isinstance(v, float):
        return round(v, 4)
    if isinstance(v, str):
        return v.strip() or None
    return v


def read_sheet(ws_values, ws_links, cols, link_col):
    # Values come from the cached formula results; links from the formula workbook.
    rows = []
    links = ws_links.iter_rows(min_row=2)
    for cells, link_cells in zip(ws_values.iter_rows(min_row=2), links):
        if all(c.value is None for c in cells):
            continue
        rec = {k: clean(c.value) for k, c in zip(cols, cells)}
        link = link_cells[cols.index(link_col)].hyperlink
        rec[link_col] = link.target if link else None
        rows.append(rec)
    return rows


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    wb_links = openpyxl.load_workbook(xlsx)
    cents = centroids()

    ops = read_sheet(wb["Oportunidades"], wb_links["Oportunidades"], OPORTUNIDADES_COLS, "url")
    missing = set()
    for o in ops:
        c = cents.get(norm(o["barrio"]))
        if not c:
            missing.add(o["barrio"])
            continue
        o["lat"], o["lon"] = jitter(f'{o["id"]}|{o["direccion"]}', *c)
        o["approx"] = True
    if missing:
        sys.exit(f"Sin centroide para: {sorted(missing)}. Agregalos a EXTRA_CENTROIDS.")

    avisos = read_sheet(wb["Todos los avisos"], wb_links["Todos los avisos"], AVISOS_COLS, "url")
    reciclado = wb["Referencias"]["B2"].value

    meta = {"source": Path(xlsx).name, "imported": date.today().isoformat(), "costo_reciclado_m2": reciclado}
    (SITE_DATA / "oportunidades.json").write_text(json.dumps({**meta, "items": ops}, ensure_ascii=False, separators=(",", ":")))
    (SITE_DATA / "avisos.json").write_text(json.dumps({**meta, "cols": AVISOS_COLS, "rows": [[a[c] for c in AVISOS_COLS] for a in avisos]}, ensure_ascii=False, separators=(",", ":")))
    print(f"{len(ops)} oportunidades, {len(avisos)} avisos")


if __name__ == "__main__":
    main(sys.argv[1])
